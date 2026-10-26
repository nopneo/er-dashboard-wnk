"""รายงาน "พยาบาลตรวจแทนแพทย์" ช่วง 16:00–20:00 น. (ตามที่หัวหน้างาน ER ขอ)

นิยาม (ตรวจกับข้อมูล รพ.วังน้ำเขียว ก.ค.–ต.ค. 2569):
  - ผู้ป่วยในช่วง = er_regist ที่ enter_er_time อยู่ 16:00:00–20:00:59 (ตั้งด้วย ER_NP_START / ER_NP_END)
  - ผู้ตรวจ = ผู้ลงโรคหลัก (ovstdiag diagtype 1) ถ้าว่างใช้ ovst.doctor
      พยาบาล = doctor.provider_type_code 03 (ตั้งด้วย ER_NURSE_PROVIDER_TYPES)
      แพทย์ = provider_type_code 01 หรือ position_id 1 (บางรายไม่ได้ลง provider_type_code)
  - "ตรวจแทนแพทย์" = พยาบาลเป็นผู้ตรวจ และไม่ใช่งานบริการ ได้แก่ ผู้ป่วยรับบริการอื่น ๆ (er_pt_type 4 = ทำแผล ฉีดยา)
    หรือโรคหลัก Z48 (ทำแผล ตัดไหม) / Z23–Z27 (วัคซีน) ที่ ER ลงเป็นผู้ป่วยทั่วไป
    — แยกนับเป็น "พยาบาลให้บริการ" ไม่ปนกับการตรวจ
  - ปี = ปีงบประมาณ (ต.ค.–ก.ย.) · เฉลี่ยต่อเดือนนับเฉพาะเดือนที่มีข้อมูลครบทั้งเดือน

ดึงเป็นรายครั้ง (ไม่มี HN / ชื่อผู้ป่วย) แล้วรวมใน Python — ทั้งหน้าเว็บและ Excel ใช้ชุดเดียวกัน
"""

import os
import re
import unicodedata
from calendar import monthrange
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO

from cache import TTLCache
from db import fetch, num
from queries import CACHE_TTL, SCREENED, build_slot, DISPOSITION, LEVEL, LEVEL_COLOR, LEVEL_NAME, TH_MONTH, _pct, _thai_date, _thai_datetime


def _hhmm(name, default):
    v = os.getenv(name, default)
    return v if re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", v) else default


START = _hhmm("ER_NP_START", "16:00")
END = _hhmm("ER_NP_END", "20:00")
NURSE_TYPES = tuple(c.strip() for c in os.getenv("ER_NURSE_PROVIDER_TYPES", "03").split(",")
                    if re.fullmatch(r"\d{1,3}", c.strip())) or ("03",)
#: ช่วงยาวสุดของรายงานนี้ — query เบากว่าแดชบอร์ดหลักมาก จึงดูย้อนหลังได้หลายปี
MAX_DAYS = int(os.getenv("ER_NP_MAX_DAYS", "1827"))

#: er_pt_type ของ HOSxP
PT_TYPE = {1: "ฉุกเฉิน", 2: "อุบัติเหตุ", 3: "ตรวจโรคทั่วไป", 4: "รับบริการอื่น ๆ", 5: "อุบัติเหตุจราจร",
           6: "ชันสูตรศพ"}
SVC_TYPES = (4, 6)
#: โรคหลักที่เป็นงานบริการแม้ลงประเภทผู้ป่วยอื่น — Z48 ทำแผล/ตัดไหม, Z23–Z27 วัคซีน (รหัสใน HOSxP ไม่มีจุด)
SVC_PDX = re.compile(r"Z(48|2[3-7])")
TH_WEEKDAY = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]

GROUPS = [  # (key, ชื่อ, สี)
    ("nurse", "พยาบาลตรวจแทนแพทย์", "#db2777"),
    ("doctor", "แพทย์ตรวจ", "#3b82f6"),
    ("svc", "พยาบาลให้บริการ (ทำแผล/ฉีดยา ฯลฯ)", "#f59e0b"),
    ("other", "อื่น ๆ / ไม่ระบุผู้ตรวจ", "#94a3b8"),
]
GROUP_NAME = {k: n for k, n, _ in GROUPS}

_cache = TTLCache(maxsize=16)


#: คำนำหน้าที่เขียนแยกจากชื่อด้วยช่องว่าง (เช่น "นาง สมศรี ใจดี") — ต้องเก็บไว้คู่กับชื่อ
TITLES = {"นาย", "นาง", "นางสาว", "น.ส.", "ว่าที่", "พว.", "พย.", "พยาบาลวิชาชีพ"}


#: จำนวนตัวอักษรต้นนามสกุลที่แสดง (ไม่นับสระ/วรรณยุกต์บนล่าง)
SURNAME_SHOW = 2


def _mask_surname(name):
    """คำนำหน้า + ชื่อ + นามสกุลบางส่วน เช่น "นางสาวสมศรี ใจดี" → "นางสาวสมศรี ใจxxx"
    ใช้ทั้งหน้าเว็บ, Excel และไฟล์ส่งออก — นามสกุลเต็มไม่ออกจากฟังก์ชันนี้ (นามสกุลสั้นแสดงน้อยลง ไม่ครบคำ)"""
    p = name.split()
    i = 0
    while i < len(p) - 1 and (p[i] in TITLES or p[i].endswith(".")):
        i += 1
    first, surname = " ".join(p[:i + 1]), "".join(p[i + 1:])
    if not surname:
        return first
    show = min(SURNAME_SHOW, sum(unicodedata.category(c) != "Mn" for c in surname) - 1)
    keep, base = "", 0
    for ch in surname:
        if unicodedata.category(ch) != "Mn":      # สระ/วรรณยุกต์บนล่างติดไปกับตัวหน้า
            if base == show:
                break
            base += 1
        keep += ch
    return f"{first} {keep}xxx"


def _fy(d):
    """ปีงบประมาณ (พ.ศ.) ของวันที่"""
    return d.year + 543 + (1 if d.month >= 10 else 0)


def _month_label(ym):
    y, m = ym
    return f"{TH_MONTH[m - 1]} {str(y + 543)[2:]}"


def _rows(d1, d2):
    nt = ", ".join(f"'{c}'" for c in NURSE_TYPES)
    end = (datetime.strptime(END, "%H:%M") + timedelta(minutes=1)).strftime("%H:%M")
    return fetch(f"""
        select t.vstdate, t.tm, t.pt, t.lvl, t.dch, t.pdx,
          case when ex.provider_type_code in ({nt}) then 'N'
               when ex.provider_type_code = '01' or ex.position_id = 1 then 'D' else 'O' end ex,
          ex.code exc, ex.name exn
        from (
          select e.vstdate, time_format(e.enter_er_time, '%H:%i') tm, ifnull(e.er_pt_type, 0) pt,
            {LEVEL} lvl, ifnull(e.er_dch_type, 0) dch,
            (select x.icd10 from ovstdiag x where x.vn = e.vn and x.diagtype = '1' limit 1) pdx,
            coalesce((select nullif(x.doctor, '') from ovstdiag x
                      where x.vn = e.vn and x.diagtype = '1' limit 1), o.doctor) exam
          from er_regist e join ovst o on o.vn = e.vn
            left join er_emergency_type et on et.er_emergency_type = e.er_emergency_type
          where e.vstdate between '{d1}' and '{d2}' and e.enter_er_time is not null and {SCREENED}
            and time(e.enter_er_time) >= '{START}:00' and time(e.enter_er_time) < '{end}:00'
        ) t left join doctor ex on ex.code = t.exam
        order by t.vstdate, t.tm""")


def _holidays(d1, d2):
    """วันหยุดตามปฏิทิน HOSxP (ตาราง holiday รวมเสาร์-อาทิตย์แล้ว)"""
    return {str(r["d"])[:10] for r in fetch(f"select holiday_date d from holiday "
                                            f"where holiday_date between '{d1}' and '{d2}'")}


def _icd_names(codes):
    codes = sorted({c for c in codes if c and re.fullmatch(r"[A-Z0-9.]{2,9}", c)})
    if not codes:
        return {}
    out = {}
    # ทีละ 150 รหัส — คำสั่งยาวเกิน ~4 KB ผ่าน ssh บน Windows แล้วถูกตัด (โหมด ssh)
    for i in range(0, len(codes), 150):
        part = ", ".join(f"'{c}'" for c in codes[i:i + 150])
        for r in fetch(f"select code, ifnull(nullif(tname, ''), name) name from icd101 where code in ({part})"):
            if r.get("code"):     # ชื่อโรคบางรายการมีขึ้นบรรทัดใหม่ — โหมด ssh ได้แถวเศษ
                out[r["code"]] = (r.get("name") or r["code"]).strip()
    return out


def _group(r):
    if r["ex"] == "N":
        return "svc" if r["pt"] in SVC_TYPES or SVC_PDX.match(r["pdx"].upper()) else "nurse"
    return "doctor" if r["ex"] == "D" else "other"


def _dispo_label(dch):
    for _, label, codes, _ in DISPOSITION:
        if dch in codes:
            return label
    return "อื่น ๆ / ยังไม่ระบุ"


def _counts(rows):
    c = Counter(r["g"] for r in rows)
    out = {k: c.get(k, 0) for k, _, _ in GROUPS}
    out["total"] = len(rows)
    out["pct"] = _pct(out["nurse"], out["total"])
    return out


def build(d1, d2):
    key = (d1, d2)
    hit = _cache.get(key, CACHE_TTL or 1)
    if hit is not None:
        return hit
    with build_slot():
        hit = _cache.get(key, CACHE_TTL or 1)
        return hit if hit is not None else _build(d1, d2, key)


def _build(d1, d2, key):
    a, b = date.fromisoformat(d1), date.fromisoformat(d2)
    rows = []
    for r in _rows(d1, d2):
        r = {"d": date.fromisoformat(str(r["vstdate"])[:10]), "tm": str(r["tm"]), "pt": int(num(r["pt"])),
             "lvl": int(num(r["lvl"], 5)), "dch": int(num(r["dch"])), "pdx": r["pdx"] or "",
             "ex": r["ex"], "exc": r["exc"] or "", "exn": _mask_surname(r["exn"] or "")}
        r["g"] = _group(r)
        rows.append(r)
    nurse = [r for r in rows if r["g"] == "nurse"]
    days = (b - a).days + 1
    hol = _holidays(d1, d2)
    for r in rows:
        r["hol"] = r["d"].isoformat() in hol or r["d"].weekday() >= 5

    # ---- รายวัน (ทุกวันในช่วง แม้ไม่มีผู้ป่วย — ใช้ต่อใน Excel)
    by_day = defaultdict(list)
    for r in rows:
        by_day[r["d"]].append(r)
    daily = []
    for i in range(days):
        d = a + timedelta(days=i)
        c = _counts(by_day.get(d, []))
        c.update({"date": d, "weekday": TH_WEEKDAY[d.weekday()],
                  "holiday": d.isoformat() in hol or d.weekday() >= 5})
        daily.append(c)

    # ---- รายเดือน
    by_month = defaultdict(list)
    for dd in daily:
        by_month[(dd["date"].year, dd["date"].month)].append(dd)
    monthly = []
    for ym in sorted(by_month):
        ds = by_month[ym]
        c = {k: sum(x[k] for x in ds) for k in ("total", "nurse", "doctor", "svc", "other")}
        full = len(ds) == monthrange(*ym)[1]
        c.update({"ym": ym, "label": _month_label(ym), "fy": _fy(date(ym[0], ym[1], 1)), "days": len(ds),
                  "full": full, "pct": _pct(c["nurse"], c["total"]),
                  "per_day": round(c["nurse"] / len(ds), 1)})
        monthly.append(c)

    # ---- รายปีงบประมาณ
    yearly = []
    for fy in sorted({m["fy"] for m in monthly}):
        ms = [m for m in monthly if m["fy"] == fy]
        full = [m for m in ms if m["full"]]
        c = {k: sum(m[k] for m in ms) for k in ("total", "nurse", "doctor", "svc", "other")}
        nd = sum(m["days"] for m in ms)
        c.update({"fy": fy, "months": len(ms), "full_months": len(full), "days": nd,
                  "complete": nd == (366 if (fy - 543) % 4 == 0 else 365),
                  "pct": _pct(c["nurse"], c["total"]), "per_day": round(c["nurse"] / nd, 1) if nd else 0,
                  "per_month": round(sum(m["nurse"] for m in full) / len(full), 1) if full else None,
                  "max": max(full, key=lambda m: m["nurse"]) if full else None,
                  "min": min(full, key=lambda m: m["nurse"]) if full else None})
        c["max_label"] = c["max"]["label"] if c["max"] else "–"
        c["min_label"] = c["min"]["label"] if c["min"] else "–"
        yearly.append(c)

    tot = _counts(rows)
    full_m = [m for m in monthly if m["full"]]
    kpi = dict(tot, days=days, per_day=round(tot["nurse"] / days, 1) if days else 0,
               per_month=round(sum(m["nurse"] for m in full_m) / len(full_m), 1) if full_m else None,
               full_months=len(full_m),
               nurse_l12=sum(r["lvl"] <= 2 for r in nurse),
               nurse_ar=sum(r["dch"] in (2, 3) for r in nurse),
               # เฉพาะเคสที่มีการตรวจ (พยาบาลตรวจแทนแพทย์ + แพทย์ตรวจ) ไม่นับงานบริการและไม่ระบุผู้ตรวจ
               l345=sum(3 <= r["lvl"] <= 5 and r["g"] in ("nurse", "doctor") for r in rows))
    kpi["l345_per_day"] = round(kpi["l345"] / days, 1) if days else 0
    kpi["nurse_l12_pct"] = _pct(kpi["nurse_l12"], tot["nurse"])
    kpi["nurse_ar_pct"] = _pct(kpi["nurse_ar"], tot["nurse"])

    # ---- รายชั่วโมง / วันในสัปดาห์ / วันทำการ-วันหยุด
    hours = []
    for h in range(int(START[:2]), int(END[:2]) + 1):
        hr = [r for r in rows if int(r["tm"][:2]) == h]
        c = _counts(hr)
        last = END if h == int(END[:2]) else f"{h:02d}:59"
        c.update({"h": h, "label": f"{h:02d}:00" if last == f"{h:02d}:00" else f"{h:02d}:00–{last}"})
        hours.append(c)
    weekdays = []
    for i, name in enumerate(TH_WEEKDAY):
        ds = [x for x in daily if x["date"].weekday() == i]
        n = sum(x["nurse"] for x in ds)
        weekdays.append({"label": name, "days": len(ds), "nurse": n, "total": sum(x["total"] for x in ds),
                         "avg": round(n / len(ds), 1) if ds else 0})
    daytype = []
    for label, flag in (("วันทำการ", False), ("วันหยุด / เสาร์-อาทิตย์", True)):
        ds = [x for x in daily if x["holiday"] == flag]
        n = sum(x["nurse"] for x in ds)
        daytype.append({"label": label, "days": len(ds), "nurse": n, "total": sum(x["total"] for x in ds),
                        "avg": round(n / len(ds), 1) if ds else 0, "pct": _pct(n, sum(x["total"] for x in ds))})

    # ---- ลักษณะผู้ป่วยที่พยาบาลตรวจ
    nn = len(nurse)
    pt = Counter(r["pt"] for r in nurse)
    pt_colors = {1: "#ef4444", 2: "#f59e0b", 3: "#22a06b", 5: "#7c2d12"}
    pt_items = [{"label": PT_TYPE.get(k, "ไม่ระบุ"), "value": v, "pct": _pct(v, nn),
                 "color": pt_colors.get(k, "#94a3b8")} for k, v in sorted(pt.items(), key=lambda x: -x[1])]
    lv = Counter(r["lvl"] for r in nurse)
    levels = [{"level": k, "name": LEVEL_NAME[k], "color": LEVEL_COLOR[k], "value": lv.get(k, 0),
               "pct": _pct(lv.get(k, 0), nn)} for k in range(1, 6)]
    dp = Counter(_dispo_label(r["dch"]) for r in nurse)
    dispo = []
    for _, label, _, color in DISPOSITION + [(None, "อื่น ๆ / ยังไม่ระบุ", (), "#94a3b8")]:
        if dp.get(label):
            dispo.append({"label": label, "value": dp[label], "pct": _pct(dp[label], nn), "color": color})

    names = _icd_names({r["pdx"] for r in rows})
    dx = Counter(r["pdx"] for r in nurse if r["pdx"])
    top = max(dx.values(), default=1)
    top_dx = [{"rank": i, "code": c, "label": names.get(c, c), "value": v, "pct": _pct(v, nn),
               "bar": _pct(v, top)} for i, (c, v) in enumerate(dx.most_common(10), 1)]
    no_dx = sum(1 for r in nurse if not r["pdx"])

    staff = Counter((r["exc"], r["exn"] or r["exc"]) for r in nurse)
    stop = max(staff.values(), default=1)
    staff_items = [{"code": c, "label": n, "value": v, "pct": _pct(v, nn), "bar": _pct(v, stop)}
                   for (c, n), v in staff.most_common()]

    for r in rows:
        r["dx_name"] = names.get(r["pdx"], "")

    data = {
        "d1": d1, "d2": d2, "start": START, "end": END, "max_days": MAX_DAYS,
        "period_label": f"{_thai_date(a)} - {_thai_date(b)}",
        "generated_at": _thai_datetime(), "groups": GROUPS,
        "kpi": kpi, "daily": daily, "monthly": monthly, "yearly": yearly, "hours": hours,
        "weekdays": weekdays, "daytype": daytype, "pt": pt_items, "levels": levels, "dispo": dispo,
        "top_dx": top_dx, "no_dx": no_dx, "staff": staff_items, "rows": rows,
    }
    _cache.set(key, data)
    return data


# ================================================================ Excel

def to_xlsx(d, hospital=""):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="B91C1C")
    tot_fill = PatternFill("solid", fgColor="FDE2E2")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="F3C4C4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()

    def sheet(title, header, rows, widths=None, total=None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font, c.alignment, c.border = head_fill, head_font, center, border
        for r in rows:
            ws.append(list(r))
        if total:
            ws.append(list(total))
            for c in ws[ws.max_row]:
                c.font, c.fill = Font(bold=True), tot_fill
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = border
                if isinstance(c.value, float):
                    c.number_format = "0.0"
                elif isinstance(c.value, date):
                    c.number_format = "dd/mm/yyyy"
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32
        return ws

    k = d["kpi"]
    window = f"{d['start']}–{d['end']} น."
    # ---- สรุป
    ws = wb.active
    ws.title = "สรุป"
    lines = [
        (f"รายงานพยาบาลตรวจแทนแพทย์ ช่วง {window}", None),
        (hospital, None),
        ("ช่วงข้อมูล", d["period_label"]),
        ("จำนวนวัน", k["days"]),
        ("ข้อมูล ณ", d["generated_at"]),
        (None, None),
        (f"ผู้ป่วย ER ที่มาช่วง {window} ทั้งหมด", k["total"]),
        ("พยาบาลตรวจแทนแพทย์ (ราย)", k["nurse"]),
        ("ร้อยละของผู้ป่วยในช่วงเวลา", k["pct"]),
        ("เฉลี่ยต่อวัน (ราย)", k["per_day"]),
        (f"เฉลี่ยต่อเดือน (ราย · จาก {k['full_months']} เดือนที่ครบเดือน)", k["per_month"] if k["per_month"] is not None else "–"),
        ("แพทย์ตรวจ (ราย)", k["doctor"]),
        ("พยาบาลให้บริการ ทำแผล/ฉีดยา ฯลฯ (ราย)", k["svc"]),
        ("อื่น ๆ / ไม่ระบุผู้ตรวจ (ราย)", k["other"]),
        ("พยาบาลตรวจ · ระดับ 1-2 (ราย)", k["nurse_l12"]),
        ("พยาบาลตรวจ · Admit/Refer (ราย)", k["nurse_ar"]),
        ("ระดับ 3-5 ที่พยาบาล+แพทย์ตรวจ ไม่รวมงานบริการ (ราย)", k["l345"]),
        ("ระดับ 3-5 ที่พยาบาล+แพทย์ตรวจ เฉลี่ยต่อวัน (ราย)", k["l345_per_day"]),
        (None, None),
        ("นิยาม", None),
        ("ช่วงเวลา", f"เวลาเข้า ER (enter_er_time) {d['start']}:00–{d['end']}:59 ทุกวัน"),
        ("ผู้ตรวจ", "ผู้ลงโรคหลัก (ovstdiag) ถ้าว่างใช้ผู้ตรวจของ visit (ovst.doctor) · พยาบาล = ประเภทผู้ให้บริการ 03"),
        ("ตรวจแทนแพทย์", "พยาบาลเป็นผู้ตรวจ ไม่รวมผู้ป่วยรับบริการอื่น ๆ (ทำแผล ฉีดยา) และโรคหลัก Z48 (ทำแผล ตัดไหม) / "
                         "Z23–Z27 (วัคซีน) ซึ่งแยกนับเป็น 'พยาบาลให้บริการ'"),
        ("ปี", "ปีงบประมาณ ต.ค.–ก.ย. · เฉลี่ยต่อเดือนคิดเฉพาะเดือนที่มีข้อมูลครบทั้งเดือน"),
        ("ความเป็นส่วนตัว", "ไม่มี HN / ชื่อผู้ป่วย · ใช้ภายในโรงพยาบาลเท่านั้น"),
    ]
    for a, b in lines:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=14, color="7F1D1D")
    ws["A20"].font = Font(bold=True, color="7F1D1D")
    for r in range(7, 19):
        ws.cell(r, 1).font = Font(bold=r == 8)
        ws.cell(r, 2).font = Font(bold=r == 8, color="B91C1C" if r == 8 else "000000")
    ws["B9"].number_format = '0.0"%"'
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 90

    grp_head = ["ผู้ป่วยในช่วงเวลา", "พยาบาลตรวจแทนแพทย์", "% พยาบาลตรวจ", "แพทย์ตรวจ",
                "พยาบาลให้บริการ", "อื่น ๆ/ไม่ระบุ"]

    def grp(x):
        return [x["total"], x["nurse"], x["pct"], x["doctor"], x["svc"], x["other"]]

    # ---- รายปีงบ
    sheet("รายปีงบประมาณ",
          ["ปีงบประมาณ", "จำนวนวัน", "จำนวนเดือน", "เดือนที่ครบ"] + grp_head
          + ["เฉลี่ย/เดือน (เดือนที่ครบ)", "เฉลี่ย/วัน", "เดือนสูงสุด", "เดือนต่ำสุด"],
          [[y["fy"], y["days"], y["months"], y["full_months"]] + grp(y)
           + [y["per_month"] if y["per_month"] is not None else "–", y["per_day"], y["max_label"], y["min_label"]]
           for y in d["yearly"]],
          [13, 10, 11, 11, 15, 15, 12, 11, 14, 12, 18, 11, 13, 13])

    # ---- รายเดือน + กราฟ
    ms = d["monthly"]
    ws = sheet("รายเดือน",
               ["เดือน", "ปีงบ", "เดือน (ค.ศ.)", "จำนวนวัน", "ครบเดือน"] + grp_head + ["เฉลี่ย/วัน"],
               [[m["label"], m["fy"], f"{m['ym'][0]}-{m['ym'][1]:02d}", m["days"], "ครบ" if m["full"] else "ไม่ครบ"]
                + grp(m) + [m["per_day"]] for m in ms],
               [11, 8, 12, 10, 10, 15, 15, 12, 11, 14, 12, 11],
               total=["รวม", "", "", k["days"], ""] + grp(k) + [k["per_day"]])
    if ms:
        n = len(ms) + 1
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
        ch.title = f"ผู้ป่วยช่วง {window} รายเดือน แยกผู้ตรวจ"
        ch.y_axis.title = "ราย"
        for col in (7, 9, 10):
            ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
        ch.height, ch.width = 9, max(18, len(ms) * 1.3)
        ws.add_chart(ch, "N2")

    # ---- รายวัน
    sheet("รายวัน", ["วันที่", "วัน", "วันหยุด", "ปีงบ", "เดือน"] + grp_head,
          [[x["date"], x["weekday"], "หยุด" if x["holiday"] else "", _fy(x["date"]),
            _month_label((x["date"].year, x["date"].month))] + grp(x) for x in d["daily"]],
          [12, 10, 8, 8, 10, 15, 15, 12, 11, 14, 12],
          total=["รวม", "", "", "", ""] + grp(k))
    wb["รายวัน"].auto_filter.ref = f"A1:K{len(d['daily']) + 1}"

    # ---- รายชั่วโมง / วัน
    ws = sheet("ชั่วโมง-วัน", ["ชั่วโมงที่มา ER"] + grp_head,
               [[h["label"]] + grp(h) for h in d["hours"]], [22, 15, 15, 12, 11, 14, 12])
    ws.append([])
    base = ws.max_row + 1
    ws.append(["วันในสัปดาห์", "จำนวนวัน", "พยาบาลตรวจแทนแพทย์", "เฉลี่ย/วัน", "ผู้ป่วยในช่วงเวลา"])
    for w in d["weekdays"]:
        ws.append([w["label"], w["days"], w["nurse"], w["avg"], w["total"]])
    ws.append([])
    ws.append(["ประเภทวัน", "จำนวนวัน", "พยาบาลตรวจแทนแพทย์", "เฉลี่ย/วัน", "ผู้ป่วยในช่วงเวลา", "% พยาบาลตรวจ"])
    hdr2 = ws.max_row
    for t in d["daytype"]:
        ws.append([t["label"], t["days"], t["nurse"], t["avg"], t["total"], t["pct"]])
    for r in (base, hdr2):
        for c in ws[r]:
            if c.value is not None:
                c.fill, c.font, c.alignment = head_fill, head_font, center

    # ---- ลักษณะผู้ป่วยที่พยาบาลตรวจ
    ws = sheet("ลักษณะผู้ป่วย", ["ประเภทผู้ป่วย", "ราย", "%"],
               [[x["label"], x["value"], x["pct"]] for x in d["pt"]], [30, 10, 10])
    for title, items in (("ระดับความเร่งด่วน", [{"label": f"ระดับ {x['level']} {x['name']}", **x} for x in d["levels"]]),
                         ("ผลการจำหน่าย", d["dispo"])):
        ws.append([])
        ws.append([title, "ราย", "%"])
        for c in ws[ws.max_row]:
            c.fill, c.font, c.alignment = head_fill, head_font, center
        for x in items:
            ws.append([x["label"], x["value"], x["pct"]])

    # ---- โรค
    dx = Counter(r["pdx"] for r in d["rows"] if r["g"] == "nurse" and r["pdx"])
    names = {r["pdx"]: r["dx_name"] for r in d["rows"]}
    sheet("โรค (พยาบาลตรวจ)", ["อันดับ", "ICD-10", "ชื่อโรค", "ราย", "%"],
          [[i, c, names.get(c, ""), v, _pct(v, k["nurse"])] for i, (c, v) in enumerate(dx.most_common(), 1)],
          [8, 10, 60, 8, 8])

    # ---- รายพยาบาล
    sheet("รายพยาบาล", ["รหัส", "ชื่อพยาบาลผู้ตรวจ", "ราย", "%"],
          [[s["code"], s["label"], s["value"], s["pct"]] for s in d["staff"]], [10, 36, 8, 8])

    # ---- รายการรายครั้ง (ไม่มีตัวระบุผู้ป่วย)
    ws = sheet("รายการ", ["วันที่", "เวลาเข้า ER", "วัน", "วันหยุด", "ปีงบ", "เดือน", "กลุ่ม", "ประเภทผู้ป่วย",
                         "ระดับ", "การจำหน่าย", "ICD-10", "โรคหลัก", "ผู้ตรวจ"],
               [[r["d"], r["tm"], TH_WEEKDAY[r["d"].weekday()], "หยุด" if r["hol"] else "",
                 _fy(r["d"]), _month_label((r["d"].year, r["d"].month)), GROUP_NAME[r["g"]],
                 PT_TYPE.get(r["pt"], "ไม่ระบุ"), r["lvl"], _dispo_label(r["dch"]), r["pdx"], r["dx_name"],
                 r["exn"] or r["exc"]] for r in d["rows"]],
               [12, 10, 10, 8, 8, 10, 30, 16, 7, 22, 9, 46, 30])
    ws.auto_filter.ref = f"A1:M{len(d['rows']) + 1}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
